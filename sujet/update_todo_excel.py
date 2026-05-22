from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter
import os

path = r"d:\xampp\htdocs\S6\evaluation\newApp\vue-project\sujet\Todo Cloud S5.xlsx"

# The current todo list to write (ID, Task, Status)
tasks = [
    (1, 'ETU3314 — Checklist sujet & vérif finale', 'completed'),
    (2, 'Protéger Backoffice (login/MDP)', 'completed'),
    (3, 'Bouton Réinitialiser données (Reset)', 'completed'),
    (4, 'Page Import (3 CSV + ZIP images)', 'completed'),
    (5, 'Validation import (colonnes/date/montants)', 'completed'),
    (6, 'Page Commandes (afficher + changer état)', 'completed'),
    (7, 'Boutons Livrer / Annuler (Commandes)', 'completed'),
    (8, 'Page Tableau de bord (par jour + total)', 'completed'),
    (9, 'Page Statistiques (Ventes HT, Achats HT, Bénéfice)', 'completed'),
    (10, 'Tableau stock par catégorie (physique/réservé/dispo)', 'completed'),
    (11, 'Page Gestion Stock + évolution journalière', 'completed'),
    (12, 'Afficher qtés produit sur fiche produit', 'completed'),
    (13, 'Front: page accueil + fiche produit', 'completed'),
    (14, 'Front: gestion panier & checkout (COD)', 'completed'),
    (15, 'Front: Mes Commandes + états', 'completed'),
    (16, 'Recherche multicritère (nom/catégorie/prix)', 'completed'),
    (17, 'Badges HOT / NEW (logic)', 'completed'),
    (18, 'Créer module PrestaShop `mon_module` (shiporder)', 'completed'),
    (19, 'Fichier d\'aide `docs/aide.md` (import modifs)', 'completed'),
    (20, 'Vérifier totaux commandes (cohérence)', 'not-started'),
    (21, 'Rapport commandes affectées (totaux mismatches)', 'not-started')
]

# Load or create workbook
if os.path.exists(path):
    wb = load_workbook(path)
else:
    wb = Workbook()

# Remove existing sheet if present
sheet_name = 'TodoUpdated'
if sheet_name in wb.sheetnames:
    std = wb[sheet_name]
    wb.remove(std)

ws = wb.create_sheet(sheet_name)

# Write header
ws.append(['ID', 'Tâche', 'Statut'])

for t in tasks:
    ws.append(list(t))

# Autosize columns
for col in ws.columns:
    max_length = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        try:
            value = str(cell.value)
            if len(value) > max_length:
                max_length = len(value)
        except:
            pass
    ws.column_dimensions[col_letter].width = max(15, min(60, max_length + 2))

wb.save(path)
print('Updated', path)
